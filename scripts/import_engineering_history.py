"""Import engineering visits from an Excel sheet with båt, dato and info columns."""
import argparse
import hashlib
import json
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "assets/Enginering table.xlsx"
DEFAULT_OUTPUT = ROOT / "wwwroot/data/vessels/engineering.json"
DEFAULT_NOTES = ROOT / "Features/Vessels/engineering-import-notes.json"


def text(value):
    return " ".join(unicodedata.normalize("NFC", str(value or "")).split())


def identifier(prefix, value):
    return prefix + "-" + hashlib.sha256(value.encode("utf-8")).hexdigest()[:20]


def display_name(value):
    return " ".join(word if word.isupper() else word.capitalize() for word in text(value).split())


def parse_periods(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return []
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return [{"start": value.isoformat(), "end": None}]
    value = text(value)
    # A single ISO date, or a period with two ISO dates separated by space or "til".
    match = re.fullmatch(r"(\d{4}-\d{2}-\d{2})(?:\s+(?:til\s+)?(\d{4}-\d{2}-\d{2}))?", value)
    if not match:
        raise ValueError(f"Invalid date {value!r}; use YYYY-MM-DD or YYYY-MM-DD YYYY-MM-DD")
    start = date.fromisoformat(match[1])
    end = date.fromisoformat(match[2]) if match[2] else None
    if end is not None and end < start:
        raise ValueError(f"Period ends before it starts: {value!r}")
    return [{"start": start.isoformat(), "end": end.isoformat() if end and end != start else None}]


def read_rows(source, sheet_name=None):
    # Only the import command needs openpyxl; rule tests use the standard library.
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ValueError("Install the Excel reader: python -m pip install -r scripts/requirements-engineering.txt") from error

    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        columns = None
        rows = []
        for row in sheet.iter_rows():
            if columns is None:
                labels = {text(cell.value).casefold(): index for index, cell in enumerate(row) if cell.value is not None}
                if {"båt", "dato", "info"}.issubset(labels):
                    columns = {key: labels[key] for key in ("båt", "dato", "info", "id") if key in labels}
                continue
            values = {key: row[index].value for key, index in columns.items()}
            if not any(value is not None and text(value) for value in values.values()):
                continue
            location = f"{sheet.title}!{row[columns['båt']].row}"
            if any(row[index].data_type in {"f", "e"} for index in columns.values()):
                raise ValueError(f"{location}: enter values, not formulas or Excel errors")
            rows.append(dict(values, location=location))
        if columns is None:
            raise ValueError("No header with båt, dato and info was found")
        return rows
    finally:
        workbook.close()


def expand_image_dates(rows, records):
    """Use reviewed dates from an embedded image for its undated workbook row."""
    rows = list(rows)
    for record in records:
        matching = [index for index, row in enumerate(rows)
                    if text(row.get("båt")).casefold() == text(record["boat"]).casefold()
                    and not text(row.get("dato"))]
        if len(matching) != 1 or not record["dates"]:
            raise ValueError(f"Review image dates for {record['boat']}: expected one undated workbook row")
        index = matching[0]
        original = rows[index]
        description = text(original.get("info")) or record.get("description")
        rows[index:index + 1] = [dict(original, dato=value, info=description) for value in record["dates"]]
    return rows


def import_engineering(rows, date_fixes=None):
    date_fixes = date_fixes or {}
    vessels = {}
    entry_ids = set()
    warnings = []
    for number, row in enumerate(rows, start=1):
        location = row.get("location", f"Row {number}")
        name = text(row.get("båt"))
        if not name:
            raise ValueError(f"{location}: a boat name is required")
        key = name.casefold()
        vessel_id = identifier("engineering-vessel", key)
        raw_date = row.get("dato")
        if isinstance(raw_date, str) and raw_date in date_fixes:
            warnings.append(f"{location}: date corrected from {raw_date} to {date_fixes[raw_date]}")
            raw_date = date_fixes[raw_date]
        try:
            periods = parse_periods(raw_date)
        except ValueError as error:
            raise ValueError(f"{location}: {error}") from error
        description = text(row.get("info")) or None
        if not periods:
            raise ValueError(f"{location}: dato is required for each visit")
        source_id = text(row.get("id"))
        identity = "id:" + source_id if source_id else "visit:" + json.dumps([key, periods], sort_keys=True)
        entry_id = identifier("engineering-job", identity)
        if entry_id in entry_ids:
            raise ValueError(f"{location}: duplicate visit; use distinct id values for separate jobs on the same boat and date")
        entry_ids.add(entry_id)
        vessel = vessels.setdefault(vessel_id, {
            "id": vessel_id, "name": display_name(name), "note": None, "entries": [],
        })
        vessel["entries"].append({
            "id": entry_id, "periods": periods, "tonnes": None, "tonnageKind": None,
            "includeInTotal": False, "description": description, "sourceUrls": [],
        })
        if not description:
            warnings.append(f"{location}: description not registered")
    if not vessels:
        raise ValueError("No visits found; existing history was not replaced")
    for vessel in vessels.values():
        vessel["entries"].sort(key=lambda entry: (
            entry["periods"][0]["start"], entry["id"]), reverse=True)
    data = {"schemaVersion": 1, "category": "engineering",
            "vessels": sorted(vessels.values(), key=lambda vessel: vessel["name"].casefold())}
    return data, {"vessels": len(vessels), "visits": len(entry_ids), "warnings": warnings}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", nargs="?", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--sheet", help="Sheet name; defaults to the first worksheet")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--notes", type=Path, default=DEFAULT_NOTES,
                        help="Reviewed date corrections and visits transcribed from source images")
    parser.add_argument("--date-fix", action="append", default=[], metavar="OLD=YYYY-MM-DD",
                        help="Explicitly correct a source typo without editing the workbook")
    args = parser.parse_args()
    try:
        notes = json.loads(args.notes.read_text(encoding="utf-8"))
        fixes = dict(notes.get("dateCorrections", {}))
        for correction in args.date_fix:
            old, separator, new = correction.partition("=")
            if not separator or not old or not new:
                raise ValueError("--date-fix requires OLD=YYYY-MM-DD")
            parse_periods(new)
            fixes[old] = new
        rows = expand_image_dates(read_rows(args.source, args.sheet), notes.get("imageVisits", []))
        data, report = import_engineering(rows, fixes)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    except (ValueError, OSError, KeyError) as error:
        parser.exit(1, f"Import failed: {error}\n")
    print(json.dumps(report, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
